"""체험 계정의 '파일 유형별 시험' 폴더에 들어가는 표본 파일을 만듭니다 (DemoDataSeeder 가 클래스패스에서 읽음).

결과: backend/src/main/resources/demo/samples/ (파일 이름은 ASCII, 화면에 보일 이름은 DemoDataSeeder.SAMPLE_FILES)
도구(앱 의존성 아님, 임시 가상환경에 설치해 실행): pip install pillow openpyxl python-pptx pypdf
  python3 -m venv /tmp/venv && /tmp/venv/bin/pip install pillow openpyxl python-pptx pypdf && /tmp/venv/bin/python qa/scripts/demo-sample-files.py

공개 저장소에 들어가므로 글꼴을 파일에 심지 않습니다: 이미지는 Pillow 에 들어 있는 글꼴(OFL)로 영문만,
PDF 는 PDF 표준 한국어 CID 글꼴(HYGoThic-Medium, Adobe-Korea1)을 이름으로만 참조합니다(보는 쪽 글꼴로 그림).
"""
import io, json, os, zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfReader

OUT = Path(__file__).resolve().parents[2] / 'backend/src/main/resources/demo/samples'
OUT.mkdir(parents=True, exist_ok=True)


def write(name, data):
    (OUT / name).write_bytes(data if isinstance(data, bytes) else data.encode('utf-8'))


# ── 텍스트 ──
write('memo.txt', 'SmartCollab 시험용 텍스트 파일입니다.\n편집기에서 고쳐 저장하면 새 버전이 생기고, 버전 기록에서 되돌릴 수 있습니다.\n')
write('meeting.md', """# 3분기 제품 회의록

## 참석
김하늘(팀장), 이도윤, 박서연

## 결정 사항
- 공유 링크에 비밀번호와 다운로드 횟수 제한을 둔다.
- 버전 기록 화면에서 서명 상태를 한눈에 보이게 한다.
- 팀 채팅으로 보낸 파일은 드라이브와 같은 미리보기로 연다.

## 할 일
| 담당 | 내용 | 기한 |
|---|---|---|
| 이도윤 | 공유 링크 만료 알림 | 다음 회의 전 |
| 박서연 | 미리보기 화면 시안 | 이번 주 |
""")
write('server.log', ''.join(f'2026-10-12T03:{i:02d}:00Z INFO  request id={1000 + i} path=/api/folders status=200 ms={12 + i % 7}\n' for i in range(30))
      + '2026-10-12T03:30:00Z WARN  upload rejected: storage quota exceeded user=demo2\n')
write('requirements.csv', '항목,담당,상태,메모\n공유 링크 비밀번호,김하늘,완료,4자 이상\n버전 서명,이도윤,진행 중,현재 버전만\n채팅 파일 미리보기,박서연,대기,"이미지, PDF 우선"\n')
write('settings.json', json.dumps({'app': 'SmartCollab', 'features': {'share': True, 'versions': True, 'chat': True},
                                   'limits': {'uploadMb': 200, 'trashDays': 30}, 'locale': 'ko-KR'}, ensure_ascii=False, indent=2) + '\n')


# ── Word(docx): 본문만 있는 최소 문서 ──
def docx(paragraphs):
    esc = lambda s: s.replace('&', '&amp;').replace('<', '&lt;')
    body = ''.join(f'<w:p><w:r><w:t xml:space="preserve">{esc(p)}</w:t></w:r></w:p>' for p in paragraphs)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr('word/document.xml', f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{body}</w:body></w:document>')
    return buf.getvalue()


write('onboarding.docx', docx([
    '신규 입사자 안내',
    '첫 주에는 팀 드라이브의 "기획" 폴더에서 프로젝트 개요와 요구사항 체크리스트를 읽어 주세요.',
    '문서를 고칠 때는 편집기에서 저장하면 새 버전이 생기며, 다른 사람이 먼저 저장했다면 충돌 안내가 나옵니다.',
    '외부에 파일을 보낼 때는 공유 링크에 비밀번호와 만료 기간을 꼭 설정합니다.',
    '궁금한 점은 팀 채팅에 남겨 주세요.',
]))


# ── PDF: 글꼴을 심지 않고 표준 한국어 CID 글꼴을 참조 ──
def pdf(pages):
    objs = []
    def add(body):
        objs.append(body)
        return len(objs)
    font_desc = add('<< /Type /FontDescriptor /FontName /HYGoThic-Medium /Flags 4 /FontBBox [-6 -145 1003 880] '
                    '/ItalicAngle 0 /Ascent 880 /Descent -120 /CapHeight 880 /StemV 93 >>')
    cid = add(f'<< /Type /Font /Subtype /CIDFontType0 /BaseFont /HYGoThic-Medium '
              f'/CIDSystemInfo << /Registry (Adobe) /Ordering (Korea1) /Supplement 1 >> /FontDescriptor {font_desc} 0 R /DW 1000 /W [1 100 500] >>')   # 영문·숫자·공백(CID 1~100)은 반각
    font = add(f'<< /Type /Font /Subtype /Type0 /BaseFont /HYGoThic-Medium-UniKS-UCS2-H /Encoding /UniKS-UCS2-H /DescendantFonts [{cid} 0 R] >>')
    pages_id = len(objs) + 1 + 2 * len(pages)   # 페이지·내용 객체 뒤
    page_ids = []
    for lines in pages:
        ops = ['BT', '/F1 20 Tf', '72 770 Td']
        for i, (size, text) in enumerate(lines):
            if i:
                ops.append(f'0 -{size + 14} Td')
            ops.append(f'/F1 {size} Tf <{text.encode("utf-16-be").hex()}> Tj')
        ops.append('ET')
        stream = '\n'.join(ops).encode()
        content = add(f'<< /Length {len(stream)} >>\nstream\n'.encode() + stream + b'\nendstream')
        page_ids.append(add(f'<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 595 842] '
                            f'/Resources << /Font << /F1 {font} 0 R >> >> /Contents {content} 0 R >>'))
    assert add(f'<< /Type /Pages /Kids [{" ".join(f"{p} 0 R" for p in page_ids)}] /Count {len(page_ids)} >>') == pages_id
    catalog = add(f'<< /Type /Catalog /Pages {pages_id} 0 R >>')
    out = bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f'{i} 0 obj\n'.encode() + (body if isinstance(body, bytes) else body.encode()) + b'\nendobj\n'
    xref = len(out)
    out += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode()
    out += ''.join(f'{o:010d} 00000 n \n' for o in offsets).encode()
    out += f'trailer\n<< /Size {len(objs) + 1} /Root {catalog} 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    return bytes(out)


write('report.pdf', pdf([
    [(22, '3분기 서비스 운영 보고서'), (12, 'SmartCollab 팀 드라이브 시험용 문서 (2쪽)'),
     (14, '1. 요약'), (12, '이번 분기 활성 팀은 지난 분기보다 늘었고,'), (12, '공유 링크 사용이 가장 크게 증가했습니다.'),
     (12, '대부분의 증가는 외부 협력사와 파일을 주고받는 팀에서 나왔습니다.'),
     (14, '2. 저장 공간'), (12, '팀당 평균 사용량은 1.2GB 이며, 옛 버전이 전체의 30%를 차지합니다.')],
    [(14, '3. 다음 분기 계획'), (12, '- 버전 기록에서 서명 상태를 한눈에 보이게 합니다.'),
     (12, '- 휴지통 보관 기간 안내를 화면에 더합니다.'), (12, '- 팀 채팅으로 보낸 파일의 미리보기를 개선합니다.')],
]))
assert '운영 보고서' in PdfReader(OUT / 'report.pdf').pages[0].extract_text(), 'PDF 글자를 뽑을 수 있어야 함(요약·번역)'


# ── Excel·PowerPoint ──
wb = Workbook()
ws = wb.active
ws.title = '예산'
ws.append(['항목', '담당', '1분기', '2분기', '합계'])
for row in [('서버', '김하늘', 120, 130), ('저장소', '이도윤', 40, 45), ('AI 요약', '박서연', 15, 22), ('디자인', '김하늘', 60, 30)]:
    r = ws.max_row + 1
    ws.append([*row, f'=C{r}+D{r}'])
ws.append(['합계', '', '=SUM(C2:C5)', '=SUM(D2:D5)', '=SUM(E2:E5)'])
wb.create_sheet('메모').append(['단위: 만 원'])
buf = io.BytesIO(); wb.save(buf); write('budget.xlsx', buf.getvalue())

p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = 'SmartCollab 3분기 발표'; s.placeholders[1].text = '시험용 프레젠테이션'
s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = '주요 일정'
tf = s.placeholders[1].text_frame; tf.text = '공유 링크 비밀번호 기능'
for t in ['버전 기록·서명 화면 개선', '팀 채팅 파일 미리보기']:
    tf.add_paragraph().text = t
buf = io.BytesIO(); p.save(buf); write('slides.pptx', buf.getvalue())


# ── 이미지: Pillow 기본 글꼴(OFL)로 영문만 ──
font = ImageFont.load_default(28)
def image(name, fmt, color, size=(480, 300), label=None):
    im = Image.new('RGB', size, color)
    d = ImageDraw.Draw(im)
    d.rectangle([12, 12, size[0] - 12, size[1] - 12], outline='white', width=4)
    d.text((36, size[1] // 2 - 30), f'SmartCollab sample\n{label or fmt}', fill='white', font=font)
    if fmt == 'GIF':
        second = Image.new('RGB', size, '#334155')
        ImageDraw.Draw(second).text((36, size[1] // 2 - 30), 'SmartCollab sample\nGIF frame 2', fill='white', font=font)
        buf = io.BytesIO(); im.save(buf, 'GIF', save_all=True, append_images=[second], duration=800, loop=0)
    else:
        buf = io.BytesIO(); im.save(buf, fmt, **({'quality': 85} if fmt == 'JPEG' else {}))
    write(name, buf.getvalue())

image('photo.png', 'PNG', '#2563eb')
image('photo.jpg', 'JPEG', '#16a34a', label='JPG')
image('animation.gif', 'GIF', '#9333ea', label='GIF frame 1')
image('photo.webp', 'WEBP', '#ea580c')
image('photo.bmp', 'BMP', '#dc2626', size=(240, 150))

# ── 미리보기가 없는 형식 ──
buf = io.BytesIO()
with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('안내.txt', '압축 파일 안의 글입니다. 미리보기가 없는 형식(내려받기만) 확인용.\n')
    z.writestr('하위/데이터.csv', '이름,값\n가,1\n나,2\n')
write('bundle.zip', buf.getvalue())
write('data.bin', bytes(range(256)) * 8)

for f in sorted(OUT.iterdir()):
    print(f'{f.name:16} {f.stat().st_size:>8,} B')
